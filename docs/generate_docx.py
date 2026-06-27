#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Конвертация ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА.md в DOCX с оформлением по ГОСТ."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

DOCS_DIR = Path(__file__).resolve().parent
MD_PATH = DOCS_DIR / "ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА.md"
OUT_DIR = DOCS_DIR / "docx"
OUT_MAIN = OUT_DIR / "ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА.docx"
OUT_COPY = OUT_DIR / "ПЗ_Рашевский.docx"

STUDENT = "Рашевский Р.Р."
GROUP = "ПИЖ-б-о-24-1"
SUPERVISOR = "Самойлов Ф.В."
DISCIPLINE = "Технология разработки программного обеспечения"
TOPIC = "Интернет-магазин свечей ручной работы «Candels»"
YEAR = "2026"

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(14)
HEADING_SIZES = {1: Pt(16), 2: Pt(15), 3: Pt(14), 4: Pt(14)}


def _set_run_font(run, bold: bool = False, size=None) -> None:
    run.font.name = FONT_NAME
    run.font.size = size or FONT_SIZE
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)


def _set_paragraph_gost(paragraph, first_indent: bool = True) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(0)
    if first_indent:
        fmt.first_line_indent = Mm(12.5)


def _add_page_number_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_sep)
    run._r.append(fld_char_end)
    _set_run_font(run)


def _configure_section(section) -> None:
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    _add_page_number_footer(section)


def _add_centered_lines(doc: Document, lines: list[str], bold_indices: set[int] | None = None) -> None:
    bold_indices = bold_indices or set()
    for i, line in enumerate(lines):
        if not line:
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        _set_run_font(run, bold=(i in bold_indices))
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def add_title_page(doc: Document) -> None:
    ministry = [
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ",
        "РОССИЙСКОЙ ФЕДЕРАЦИИ",
        "",
        "Федеральное государственное автономное",
        "образовательное учреждение высшего образования",
        "«СЕВЕРО-КАВКАЗСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ»",
        "",
        "Институт перспективной инженерии",
        "Кафедра межинститутская базовая",
        "",
        "",
    ]
    _add_centered_lines(doc, ministry)

    for _ in range(4):
        doc.add_paragraph()

    _add_centered_lines(
        doc,
        [
            "ПОЯСНИТЕЛЬНАЯ ЗАПИСКА",
            f"к курсовому проекту по дисциплине «{DISCIPLINE}»",
            "",
            f"Тема: {TOPIC}",
        ],
        bold_indices={0},
    )

    for _ in range(6):
        doc.add_paragraph()

    info_lines = [
        f"Выполнил: студент группы {GROUP}",
        STUDENT,
        "",
        f"Научный руководитель: {SUPERVISOR}",
        "",
        "",
        f"Ставрополь, {YEAR}",
    ]
    for line in info_lines:
        p = doc.add_paragraph()
        if line:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(line)
            _set_run_font(run)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    doc.add_page_break()


def _strip_md_inline(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.strip()


def _add_text_paragraph(doc: Document, text: str, first_indent: bool = True) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(_strip_md_inline(text))
    _set_run_font(run)
    _set_paragraph_gost(p, first_indent=first_indent)


def _add_heading(doc: Document, level: int, text: str, *, page_break: bool = False) -> None:
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(_strip_md_inline(text).upper() if level == 1 else _strip_md_inline(text))
    _set_run_font(run, bold=True, size=HEADING_SIZES.get(level, FONT_SIZE))
    fmt = p.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.first_line_indent = Mm(0)
    fmt.space_before = Pt(12 if level <= 2 else 6)
    fmt.space_after = Pt(6)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    headers = rows[0]
    body = rows[1:]
    table = doc.add_table(rows=1 + len(body), cols=len(headers))
    table.style = "Table Grid"
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(_strip_md_inline(header))
        _set_run_font(run, bold=True)
    for r, row in enumerate(body):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(_strip_md_inline(val))
            _set_run_font(run)
    doc.add_paragraph()


def _add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(line.rstrip())
        _set_run_font(run, size=Pt(12))
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.left_indent = Mm(10)


def _parse_table_block(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            break
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.match(r"^[-:\s]+$", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def _is_table_separator(line: str) -> bool:
    return bool(re.match(r"^\|\s*[-:\s|]+\|\s*$", line.strip()))


def parse_markdown_to_docx(doc: Document, md_text: str) -> None:
    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    list_buf: list[str] = []

    def flush_list() -> None:
        nonlocal list_buf
        for item in list_buf:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(_strip_md_inline(item))
            _set_run_font(run)
            _set_paragraph_gost(p, first_indent=False)
        list_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                _add_code_block(doc, code_buf)
                code_buf = []
                in_code = False
            else:
                flush_list()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if stripped == "---":
            flush_list()
            i += 1
            continue

        if stripped.startswith("#"):
            flush_list()
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped.lstrip("#").strip()
            if title.upper().startswith("ПОЯСНИТЕЛЬНАЯ ЗАПИСКА") and i < 30:
                i += 1
                continue
            lvl = min(level, 4)
            need_break = lvl == 1 or (lvl == 2 and re.match(r"^\d+\.", title))
            _add_heading(doc, lvl, title, page_break=need_break)
            i += 1
            continue

        if stripped.startswith("|") and not _is_table_separator(stripped):
            flush_list()
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _is_table_separator(lines[i]):
                    table_lines.append(lines[i])
                i += 1
            rows = _parse_table_block(table_lines)
            if rows:
                _add_table(doc, rows)
            continue

        if re.match(r"^[-*]\s+", stripped):
            list_buf.append(re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            flush_list()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(_strip_md_inline(stripped))
            _set_run_font(run)
            _set_paragraph_gost(p)
            i += 1
            continue

        if not stripped:
            flush_list()
            i += 1
            continue

        flush_list()
        _add_text_paragraph(doc, stripped)
        i += 1

    flush_list()
    if in_code and code_buf:
        _add_code_block(doc, code_buf)


def build_document() -> Document:
    if not MD_PATH.exists():
        raise FileNotFoundError(
            f"Не найден {MD_PATH}. Сначала выполните: python docs/build_full_report.py"
        )
    doc = Document()
    section = doc.sections[0]
    _configure_section(section)
    add_title_page(doc)
    md_text = MD_PATH.read_text(encoding="utf-8")
    parse_markdown_to_docx(doc, md_text)
    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUT_MAIN)
    shutil.copy2(OUT_MAIN, OUT_COPY)
    print(f"DOCX saved: {OUT_MAIN}")
    print(f"Copy saved: {OUT_COPY}")


if __name__ == "__main__":
    main()

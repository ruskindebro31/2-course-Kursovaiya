import sys
sys.path.append('/root/.local/lib/python3.12/site-packages')

from docx import Document

doc = Document()

# Титульный лист
doc.add_heading('МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ', 0)
doc.add_heading('Федеральное государственное автономное образовательное учреждение высшего образования', level=2)
doc.add_heading('«СЕВЕРО-КАВКАЗСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ»', level=2)
doc.add_paragraph('\n\n\n')
doc.add_heading('Курсовой проект', level=1)
doc.add_heading('по дисциплине «Технология разработки программного обеспечения»', level=2)
doc.add_paragraph('\n\n\n')
doc.add_heading('Тема: Candels — интернет-магазин свечей', level=2)
doc.add_paragraph('\n\n\n')
doc.add_paragraph('Выполнил: студент группы ПИ-202\nИванов И.И.')
doc.add_paragraph('Проверил: доцент кафедры ПИ\nПетров П.П.')
doc.add_paragraph('\n\n\n')
doc.add_paragraph('Ставрополь — 2025', style='Normal')
doc.add_page_break()

# Содержание
doc.add_heading('Содержание', level=1)
doc.add_paragraph('1. Введение', style='Normal')
doc.add_paragraph('2. Описание проекта', style='Normal')
doc.add_paragraph('3. Требования к ПО', style='Normal')
doc.add_paragraph('4. Архитектура системы', style='Normal')
doc.add_paragraph('5. Реализация', style='Normal')
doc.add_paragraph('6. Тестирование', style='Normal')
doc.add_paragraph('7. Заключение', style='Normal')
doc.add_page_break()

# 1. Введение
doc.add_heading('1. Введение', level=1)
doc.add_paragraph(
    'Курсовой проект выполнен по дисциплине «Технология разработки программного обеспечения» '
    'в рамках траектории В: React SPA + AJAX + JWT + WebSocket. Цель проекта — разработать '
    'современное веб-приложение с использованием Django REST API и React SPA, '
    'поддерживающее аутентификацию, real-time уведомления, чат и систему отзывов.'
)

# 2. Описание проекта
doc.add_heading('2. Описание проекта', level=1)
doc.add_paragraph(
    'Проект представляет собой интернет-магазин свечей (Candels). В нём реализованы следующие функции:\n'
    '- Регистрация и аутентификация пользователей (JWT)\n'
    '- Просмотр каталога свечей с фильтрацией и пагинацией\n'
    '- Корзина и оформление заказа\n'
    '- Личный кабинет\n'
    '- Отзывы и рейтинги\n'
    '- Чат с поддержкой\n'
    '- Real-time уведомления\n'
)

# 3. Требования к ПО
doc.add_heading('3. Требования к ПО', level=1)
doc.add_paragraph(
    'Функциональные требования:\n'
    '- Авторизация и регистрация пользователей\n'
    '- Просмотр каталога с фильтрами\n'
    '- Добавление товаров в корзину\n'
    '- Оформление заказа\n'
    '- Уведомления о статусе заказа\n'
    '- Отзывы и рейтинги\n'
    '- Чат с поддержкой\n'
    '\n'
    'Нефункциональные требования:\n'
    '- Поддержка WebSocket\n'
    '- Адаптивный интерфейс\n'
    '- Безопасность (JWT, HTTPS)\n'
    '- REST API\n'
)

# 4. Архитектура системы
doc.add_heading('4. Архитектура системы', level=1)
doc.add_paragraph(
    'Система состоит из двух частей:\n'
    '1. Backend: Django REST API с JWT-аутентификацией\n'
    '2. Frontend: React SPA с использованием Vite\n'
    '\n'
    'Технологии:\n'
    '- Django, DRF, SimpleJWT\n'
    '- React, Axios, Zustand\n'
    '- WebSocket (Django Channels)\n'
    '- PostgreSQL (в будущем)\n'
)

# 5. Реализация
doc.add_heading('5. Реализация', level=1)
doc.add_paragraph(
    'Реализованы следующие модули:\n'
    '- Модели: User, Candle, Order, Cart, Review, Notification, Chat\n'
    '- API-эндпоинты: /api/candles/, /api/orders/, /api/cart/ и др.\n'
    '- JWT-аутентификация\n'
    '- Фронтенд: страницы логина, регистрации, профиля, каталога, корзины, заказов, чата\n'
    '- WebSocket-уведомления\n'
    '- Админ-панель Django с кастомными действиями\n'
)

# 6. Тестирование
doc.add_heading('6. Тестирование', level=1)
doc.add_paragraph(
    'Тестирование проводилось вручную через Swagger UI и Postman:\n'
    '- Регистрация и вход\n'
    '- Получение токена\n'
    '- Просмотр каталога\n'
    '- Добавление в корзину\n'
    '- Оформление заказа\n'
    '- Чат и уведомления\n'
    '\n'
    'Также реализованы unit-тесты для моделей и API.'
)

# 7. Заключение
doc.add_heading('7. Заключение', level=1)
doc.add_paragraph(
    'В ходе курсового проекта была разработана современная веб-система с использованием '
    'современных технологий: Django REST API, React SPA, WebSocket. Реализованы все требования '
    'траектории В, включая real-time уведомления, чат и систему отзывов. Проект готов к '
    'дальнейшему развитию и деплою.'
)

# Сохранение
doc.save('REPORT.docx')
